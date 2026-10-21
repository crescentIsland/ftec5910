from pathlib import Path
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"docs/meeting_20261021"
fig,ax=plt.subplots(figsize=(13,8))
ax.set_xlim(0,13);ax.set_ylim(0,8);ax.axis("off")
def box(x,y,w,h,title,body):
 ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle="round,pad=0.06",edgecolor="black",facecolor="#f5f5f5",linewidth=1.2))
 ax.text(x+w/2,y+h-.22,title,ha="center",va="top",fontsize=12,fontweight="bold")
 ax.text(x+w/2,y+h/2-.12,body,ha="center",va="center",fontsize=10,linespacing=1.4)
def arrow(a,b):
 ax.annotate("",xy=b,xytext=a,arrowprops={"arrowstyle":"->","color":"black","linewidth":1.3})
box(.25,6.35,3.7,1.2,"B0: Simple benchmarks","Training-price median\nDirect maximum recorded prefix bid")
box(4.65,6.35,3.7,1.2,"B1: Opening conditions","Opening price + auction duration")
box(9.1,6.35,3.65,1.2,"B2: Observed price","B1 + maximum prefix bid\n+ price rise + has-bid indicator")
arrow((3.98,6.95),(4.58,6.95));arrow((8.42,6.95),(9.03,6.95))
box(.4,3.75,3.8,1.55,"B3: B2 + Competition","Bid count; distinct bidders\nBids per bidder; concentration\nMissing-bidder count")
box(4.6,3.75,3.8,1.55,"B4: B2 + Time dynamics","Recent activity and new bidders\nTiming; concentration change\nRecent price movement")
box(8.8,3.75,3.8,1.55,"B5: B2 + Bidder feedback","Median feedback rating\nLow-rating share\nMissing-rating share")
# B3-B5 are parallel additions to B2, rather than cumulative.
ax.plot([2.3,10.95],[5.8,5.8],color="black",linewidth=1.1)
arrow((10.95,6.30),(10.95,5.8))
for x in [2.3,6.5,10.7]:arrow((x,5.8),(x,5.36))
ax.text(6.5,5.96,"Parallel tests: each group is added separately to B2",ha="center",fontsize=10)
box(.4,1.1,3.8,1.6,"B6: All candidate signals","B2 + Competition\n+ Time dynamics + Feedback")
box(4.6,1.1,3.8,1.6,"Ablation: remove one group","B6 without competition\nB6 without time dynamics\nB6 without feedback")
box(8.8,1.1,3.8,1.6,"B7: Limited interactions","B4 + new bidders x price rise\n+ concentration change x price rise\nExploratory only")
ax.text(6.5,3.15,"Follow-up checks: combined model, group removal, and limited interactions",ha="center",fontsize=10)
ax.text(6.5,.48,"Compare B3-B6 with the same-model B2 on identical auction folds; lower MAE is better.",ha="center",fontsize=11)
ax.set_title("Feature experiments: what adds value beyond observed price?",fontsize=16,fontweight="bold",pad=16)
fig.savefig(OUT/"05_experiment_groups.png",dpi=200,bbox_inches="tight");plt.close(fig)
sections=[
("1. 找到更符合研究目标的数据集",
"这段时间我筛选了几份公开数据，最后把重点放在Kaggle的Online Auctions Dataset。它记录了eBay上逐次竞价的过程，能观察信息如何在最终成交前不断形成，因此比只有静态商品属性或最终价格的表格更适合研究动态市场信号。不过它不是HarborLink的兑付数据，二者的业务机制仍有区别。",
"数据共10,681条出价、628场拍卖，包括149场Xbox、343场Palm PDA和136场Cartier手表。每行是一条出价，不是一场拍卖。排除一场起拍价不一致的PDA拍卖后，实际分析627场、10,653条出价。数据包括拍卖编号、出价金额、相对开拍时间、竞价者标识、竞价者反馈评级、起拍价、最终价格、商品类别和拍卖时长。"),
("2. 在25%、50%、75%时点分别预测",
"我把每场拍卖分别暂停在总时长的25%、50%和75%，只使用暂停前已经出现的特征，预测同一个目标：拍卖结束时的最终成交价。以7天拍卖为例，三个时点分别是第1.75、3.5和5.25天。50%作为主实验，25%和75%用于检查结果是否随阶段变化。这里的百分比是时间比例，不是训练集比例或出价条数比例。",
"每个时点，每场拍卖形成一行特征。没有出价的拍卖也保留，并用标记说明。三个截点来自同一批拍卖，不能当成三份独立数据。"),
("3. 按实验组检验增量价值",
"B0先建立两个简单基线：训练集中最终价格的中位数，以及直接使用截点前最高记录出价。B1只用起拍价和拍卖时长；B2再加入此前最高记录出价、相对起拍价涨幅以及是否已有出价。B2是后面行为特征实验的主要参照。",
"B3、B4、B5是平行实验，不是逐组累加。B3等于B2加竞争结构，B4等于B2加时间变化，B5等于B2加竞价者评级。B6把这些组全部加入；再从B6中分别去掉一组做消融。最后B7只检验两个预设交互，不穷举组合。这样能区分“新特征与价格有关”和“新特征在已有价格之外确实有用”。"),
("4. 模型、参数及数据划分",
"模型使用Ridge回归和浅层随机森林。Ridge用于检查较简单的关系，随机森林用于检查非线性和交互。本轮先固定参数，避免根据测试结果反复调整。Ridge的alpha为10；随机森林100棵树、最大深度5、最小叶样本5，随机种子5910。",
"Xbox、PDA和手表分别训练。每类商品按拍卖时长分层，做5折交叉验证，重复3次；划分种子为5910、5911、5912。同一场拍卖只能在训练或测试一边，所有实验组和时点沿用相同划分。缺失值填补与标准化只在训练折拟合。主要指标是MAE，越低越好，同时检查RMSE和R平方。文件没有真实日历开始日期，因此不能声称做了未来月份验证。"),
("5. 实验结果",
"首先，已观察价格有明确的预测价值。在50%时点，Ridge从B1到B2，Xbox的MAE由39.54降至31.19，PDA由17.39降至16.99，手表由509.28降至212.75。这里说明的是拍卖过程中形成的价格信息有用，并非所有行为特征都有效。",
"控制B2以后，竞争结构和评级没有稳定收益。时间变化组有小幅线索：Xbox随机森林从34.25降至33.35，改善2.64%；手表随机森林改善1.83%。但正向收益的探索性区间都包含零，跨模型、商品和25%/75%时点也不一致。全部加入没有普遍胜出，预设交互在50%时也未超过B4。"),
("6. 结论与下一步",
"目前可以得出的结论是：已观察价格是预测最终成交价的主要信息来源；在控制这些价格信息后，额外行为特征尚未显示稳定、明确的增量预测收益。因此，我不会把个别2%左右的改善解释为已发现显著的价格驱动因素，也不会据此证明这些行为永远无效。",
"下一步优先将B4拆成价格动量和非价格行为，再比较剩余涨价目标与更强的简单基线，并在训练集内部进行小范围调参。这样可以判断目前增益有限是信息本身不足，还是目标和模型设置限制了效果。")
]
doc=Document();sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for name,size in [("Normal",10),("Title",20),("Heading 1",12)]:
 s=doc.styles[name];s.font.name="Times New Roman";s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
doc.styles["Normal"].paragraph_format.space_after=Pt(6)
doc.add_heading("eBay拍卖实验：汇报顺序与讲稿",0)
doc.add_paragraph("2026年10月21日 | 建议5—7分钟 | 基于已完成实验")
for i,(title,p1,p2) in enumerate(sections):
 if i>0:doc.add_page_break()
 doc.add_heading(title,1)
 if i==1:doc.add_picture(str(OUT/"01_design.png"),width=Inches(6.8))
 if i==2:doc.add_picture(str(OUT/"05_experiment_groups.png"),width=Inches(6.8))
 if i==4:doc.add_picture(str(OUT/"02_price_information.png"),width=Inches(6.8))
 doc.add_paragraph(p1);doc.add_paragraph(p2)
 if i==3:
  t=doc.add_table(rows=1,cols=2);t.style="Table Grid"
  for cell,text in zip(t.rows[0].cells,["项目","关键设置"]):cell.text=text
  for a,b in [("Ridge","alpha=10"),("Random forest","n_estimators=100; max_depth=5; min_samples_leaf=5; random_state=5910"),("划分","各商品分别；按时长分层；5折 x 3次；种子5910/5911/5912"),("评估","MAE为主；RMSE、R²补充；配对误差区间为探索性")]:
   cells=t.add_row().cells;cells[0].text=a;cells[1].text=b
 if i==5:
  doc.add_paragraph("建议使用的结论句：控制已观察价格后，额外行为信号未显示稳定的增量预测收益。")
  doc.add_paragraph("避免说：所有特征都没有显著影响。B2确有预测价值；当前实验也不识别因果影响。")
for root in [doc.element,doc.styles.element]:
 for border in list(root.iter(qn("w:pBdr"))):border.getparent().remove(border)
target=OUT/"auction_meeting_speaking_notes.docx"
try:doc.save(target)
except PermissionError:
 target=OUT/"auction_meeting_speaking_notes_Updated.docx";doc.save(target)
notes="# eBay拍卖实验汇报\n\n"+"\n\n".join("## "+t+"\n\n"+a+"\n\n"+b for t,a,b in sections)
(OUT/"speaking_notes_zh.md").write_text(notes,encoding="utf-8")
print(target);print(OUT/"05_experiment_groups.png")
