from pathlib import Path
import json
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"docs/meeting_20261021"
OUT.mkdir(parents=True,exist_ok=True)
m=pd.read_csv(ROOT/"outputs/auction_feature_experiments/metrics_summary.csv")
c=pd.read_csv(ROOT/"outputs/auction_feature_experiments/paired_comparisons.csv")
names={"xbox":"Xbox","pda":"PDA","watch":"Cartier watches"}
plt.rcParams.update({"font.family":"DejaVu Sans","font.size":11,"axes.spines.top":False,"axes.spines.right":False,"axes.titleweight":"bold","savefig.facecolor":"white"})
def val(item,stage,model,exp):
 return float(m.loc[m.item_key.eq(item)&m.stage.eq(stage)&m.model.eq(model)&m.experiment.eq(exp),"mae"].iloc[0])
# Figure 1: data unit and cutoff safety.
fig,ax=plt.subplots(figsize=(12,5.8));ax.axis("off")
boxes=[(.02,.65,.25,.22,"628 auctions / 10,681 bid rows\nExclude 1 inconsistent auction\n627 auctions retained"),(.37,.65,.25,.22,"One auction = one sample\n149 Xbox | 342 PDA | 136 watches\nModels fitted separately"),(.72,.65,.25,.22,"Matched auction-level splits\n3 repeats x 5 folds\nRidge + shallow random forest")]
for x,y,w,h,text in boxes:
 ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle="round,pad=0.01",facecolor="#f2f2f2",edgecolor="black",transform=ax.transAxes))
 ax.text(x+w/2,y+h/2,text,ha="center",va="center",transform=ax.transAxes,fontsize=10)
for x in [.29,.64]:ax.annotate("",xy=(x+.065,.76),xytext=(x,.76),xycoords="axes fraction",arrowprops={"arrowstyle":"->","color":"black"})
ax.plot([.08,.92],[.35,.35],color="black",transform=ax.transAxes)
for x,t,label in [(.08,"Start","Opening conditions"),(.29,"25%","Supplementary"),(.5,"50%","Primary cutoff"),(.71,"75%","Supplementary"),(.92,"Close","Final price: label only")]:
 ax.scatter(x,.35,color="black",s=35,transform=ax.transAxes)
 ax.text(x,.43,t,ha="center",transform=ax.transAxes,fontweight="bold")
 ax.text(x,.23,label,ha="center",transform=ax.transAxes,fontsize=9)
ax.text(.5,.05,"At each cutoff: use only bids already recorded; retain auctions with no bids.",ha="center",transform=ax.transAxes)
ax.set_title("Study design: predicting the final price before the auction closes",pad=18)
fig.savefig(OUT/"01_design.png",dpi=200,bbox_inches="tight");plt.close(fig)
# Figure 2: same-model price information comparisons.
fig,axs=plt.subplots(1,3,figsize=(12,5))
for ax,item in zip(axs,names):
 a=val(item,.5,"ridge","B1_opening");b=val(item,.5,"ridge","B2_price")
 bars=ax.bar(["Opening\nconditions"," + Observed\nprice"],[a,b],color=["#c5c5c5","#555555"],width=.6)
 ax.bar_label(bars,fmt="%.2f",padding=5);ax.set_ylim(0,max(a,b)*1.25)
 ax.set(title=names[item],ylabel="MAE (recorded price units)")
 ax.text(.5,.93,f"{(a-b)/a*100:.1f}% lower MAE",ha="center",transform=ax.transAxes)
fig.suptitle("Main result: observed price adds information (50% cutoff, Ridge)",fontsize=14,fontweight="bold")
fig.tight_layout(rect=[0,0,1,.92]);fig.savefig(OUT/"02_price_information.png",dpi=200);plt.close(fig)
# Figure 3: true incremental behavior comparisons, uncertainty visible.
fig,axs=plt.subplots(1,3,figsize=(13,5.7),sharex=True)
groups=["B3_competition","B4_dynamics","B5_reputation","B6_full"]
labels=["Competition","Time dynamics","Bidder feedback","All signals"]
for ax,item in zip(axs,names):
 g=c.loc[c.item_key.eq(item)&c.stage.eq(.5)&c.model.eq("random_forest")&c.reference.eq("B2_price")].set_index("new").loc[groups]
 base=val(item,.5,"random_forest","B2_price")
 x=g.mae_gain.to_numpy()/base*100
 lo=g.ci_low.to_numpy()/base*100;hi=g.ci_high.to_numpy()/base*100
 for y,point,left,right in zip(range(4),x,lo,hi):
  ax.plot([left,right],[y,y],color="black",linewidth=1.5);ax.scatter(point,y,color="black",s=35)
  ax.text(point,y-.16,f"{point:+.2f}%",ha="center",fontsize=9)
 ax.axvline(0,color="#888888",linestyle="--")
 ax.set_yticks(range(4),labels);ax.invert_yaxis();ax.set_title(names[item])
 ax.set_xlabel("MAE reduction vs B2 (%)\nPositive = improvement")
 ax.set_xlim(-5,7)
 ax.set_ylim(3.45,-.55)
fig.suptitle("Does behavior add value beyond price? (50% cutoff, random forest)",fontsize=14,fontweight="bold")
fig.text(.5,.015,"Whiskers: exploratory paired bootstrap intervals, scaled by fixed B2 MAE. All positive-gain intervals cross zero; not confirmatory tests.",ha="center",fontsize=9)
fig.tight_layout(rect=[0,.06,1,.92]);fig.savefig(OUT/"03_behavior_increment.png",dpi=200);plt.close(fig)
# Figure 4: stage-dependence and replication.
fig,axs=plt.subplots(1,3,figsize=(12,5))
for ax,item in zip(axs,names):
 stages=[.25,.5,.75]
 b2=[val(item,s,"ridge","B2_price") for s in stages]
 b4=[val(item,s,"ridge","B4_dynamics") for s in stages]
 ax.plot([25,50,75],b2,color="black",marker="o",label="B2: observed price")
 ax.plot([25,50,75],b4,color="#888888",marker="s",linestyle="--",label="B4: + time dynamics")
 ax.set(title=names[item],xlabel="Auction duration elapsed (%)",ylabel="MAE")
 ax.set_xticks([25,50,75]);ax.grid(axis="y",alpha=.15)
 ax.legend(fontsize=8)
fig.suptitle("Stage check: price gains generally strengthen; dynamics gains do not persist (Ridge)",fontsize=13,fontweight="bold")
fig.tight_layout(rect=[0,0,1,.91]);fig.savefig(OUT/"04_stage_check.png",dpi=200);plt.close(fig)
sections=[
("开场：研究问题（约40秒）",
"我这次主要汇报公开eBay拍卖数据上的阶段性实验。我的问题不是简单地问“能不能预测最终成交价”，而是：在已经知道起拍条件和目前记录的价格以后，竞价人数、参与变化和反馈评级还能不能提供额外的预测信息？因此，重点是新增信号相对价格基线的增量价值。"),
("图1：数据和实验设计（约60秒）",
"原始数据有628场拍卖、10,681条出价记录，包含Xbox、PDA和Cartier手表。排除一场起拍价不一致的PDA拍卖后，保留627场。每场拍卖在25%、50%、75%三个时点各生成一个样本，50%是主实验。例如七天拍卖在第3.5天预测，只看此前已经出现的出价。没有出价的拍卖也保留。训练测试按整场拍卖划分，使用三次重复五折，并让所有特征组使用相同划分；不会把同一拍卖的出价同时分到训练和测试。"),
("图2：主要发现——价格信息有价值（约60秒）",
"先看同一个Ridge模型。只用起拍价和时长时，Xbox的MAE是39.54；加入此前最高记录出价等信息后降到31.19，降低21.12%。手表从509.28降到212.75，PDA从17.39降到16.99。说明拍卖过程中形成的价格信息确实有用，但不同商品的强度不同。这个结果不能直接归功于复杂模型，也不能说明额外行为特征已经有价值。三个商品的价格尺度不同，所以不能直接比较它们MAE的大小。"),
("图3：核心检验——行为是否超越价格（约80秒）",
"接下来固定B2价格基准，分别加竞争结构、时间变化、反馈评级，最后全部加入。这里展示随机森林：Xbox时间变化组的MAE降低2.64%，手表降低1.83%，但PDA没有改善。图中的点是改善比例，横线是探索性配对区间；这些正向改善对应的区间都穿过零，所以目前不能确认改善。部分负向区间不跨零，提示加特征也可能变差。Ridge的结果也不完全一致。比较准确的说法是：时间变化组值得保留为候选，但还没有稳定的增量证据；评级和累计竞争结构暂时不适合作为主要发现。"),
("图4：25%和75%的检查（约50秒）",
"最后检查这个结论是否依赖选定时点。B2在Xbox和手表中总体上随着拍卖推进而更准确，PDA变化较小。但B4在中点的小幅改善没有在另外两个时点稳定重现。这个现象说明，市场信号的价值可能依赖阶段与商品，而不是所有阶段都有效。这里仍然是相同拍卖的不同截点，不是三份独立验证数据。"),
("结尾：结论与下一步（约40秒）",
"目前的主要结论是：已观察价格是最终成交价的主要信息来源；新增行为信号的价值有限且不稳定。下一步我会把时间变化组拆成价格动量和非价格行为，研究“还会涨多少”的剩余涨价目标，并增加更强的简单基线，再在训练集内部小范围调参。我不会因为某一个组合略好就宣称发现稳定驱动因素。公开拍卖也没有HarborLink的兑付事件，所以目前贡献是动态信息评估方法，而不是对兑付机制的直接验证。")
]
qa=[
("老师问：是不是没有结果？","有两类结果：已观察价格有明确预测价值；控制价格后，附加行为信号未显示稳定增益。这回答了哪些数据值得收集、哪些复杂度暂时不必要，并非所有特征都有效。"),
("为什么只用两种模型？","先用Ridge检查较简单关系，用浅层随机森林检查非线性和交互；样本只有627场，优先限制复杂度和建立公平比较。尚不能认定换模型一定无益。"),
("为什么不按时间划分？","文件只有相对开拍时间，没有真实日历开始日期。因此采用整场拍卖交叉验证，不能声称已经验证对未来月份的预测能力。"),
("为什么不是10,681个样本？","那些是出价行，很多属于同一场拍卖。预测标签按拍卖定义，独立拍卖才是主要划分单位。"),
("评级是卖家信誉吗？","不是，是竞价者反馈评级；而且评级快照时间尚未核实，不能把它直接理解为预测时刻已经公开可得的卖家质量。"),
("2%左右算有效吗？","它是小幅候选收益，但区间跨零、跨商品和模型不一致，尚不足以确认稳定提升；实践价值还取决于误差成本和数据采集成本。"),
("为什么手表效果特别大？","手表价格范围很宽，已有出价包含很多商品差异信息。必须与当前出价等强基线比较，不能只拿起拍前弱模型来放大改善。"),
("你是否做了完整的robustness/external validation？","完成了探索性阶段比较、消融和尾部诊断；完整鲁棒性、独立外部验证和最终报告仍属于后续工作，不能将阶段性检查写成全部完成。")
]
doc=Document();sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
for name,size in [("Normal",10),("Title",20),("Heading 1",12),("Heading 2",10)]:
 s=doc.styles[name];s.font.name="Times New Roman";s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
doc.styles["Normal"].paragraph_format.space_after=Pt(6)
doc.add_heading("eBay拍卖实验汇报稿",0)
doc.add_paragraph("2026年10月21日 | 建议时长：5—6分钟 | 仅使用已完成实验")
doc.add_heading(sections[0][0],1);doc.add_paragraph(sections[0][1])
doc.add_heading("汇报顺序",1)
for t in ["问题：行为信息能否超越已有价格？","设计：627场、固定截点、整场拍卖划分。","结果：价格有用；行为有小幅线索但证据不足。","结论：按阶段与商品解释；下一步拆组和改进目标。"]:doc.add_paragraph(t,style="List Bullet")
doc.add_heading("最重要的表述",1)
doc.add_paragraph("可以说：尚未发现稳定增益。不要说：已经证明行为无效；发现因果驱动；验证了未来市场或兑付机制。")
for i,(title,speech) in enumerate(sections[1:5],1):
 doc.add_page_break();doc.add_heading(title,1)
 doc.add_picture(str(OUT/f"0{i}_{['design','price_information','behavior_increment','stage_check'][i-1]}.png"),width=Inches(6.8))
 doc.add_paragraph(speech)
 if i==3:doc.add_paragraph("读图方法：零线右边表示改善；区间跨零表示证据仍不确定。此图展示随机森林，完整报告同时包含Ridge，不应选择性只讲较好的模型。")
doc.add_page_break();doc.add_heading(sections[5][0],1);doc.add_paragraph(sections[5][1])
doc.add_heading("常见追问",1)
for q,a in qa:
 p=doc.add_paragraph();p.add_run(q+" ").bold=True;p.add_run(a)
for root in [doc.element,doc.styles.element]:
 for b in list(root.iter(qn("w:pBdr"))):b.getparent().remove(b)
doc.save(OUT/"auction_meeting_speaking_notes.docx")
text="# 2026-10-21 eBay拍卖实验汇报\n\n"
for title,speech in sections:text+="## "+title+"\n\n"+speech+"\n\n"
text+="## 常见追问\n\n"+"\n\n".join("**"+q+"**\n\n"+a for q,a in qa)
(OUT/"speaking_notes_zh.md").write_text(text,encoding="utf-8")
print(OUT)
print("Created four grayscale figures, Word speaking notes, and Markdown notes. No models rerun and no PDF generated.")
