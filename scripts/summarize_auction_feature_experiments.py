"""Create Chinese interpretation and source-register updates from experiment outputs."""
from pathlib import Path
import json
import pandas as pd
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"outputs/auction_feature_experiments"
NAMES={"xbox":"Xbox","pda":"PDA","watch":"Cartier手表"}
LABELS={"B0_median":"B0 训练成交价中位数","B0_current_bid":"补充基线 截点前最高记录出价","B1_opening":"B1 起拍条件","B2_price":"B2 ＋已观察价格","B3_competition":"B3 B2＋竞争结构","B4_dynamics":"B4 B2＋时间变化","B5_reputation":"B5 B2＋竞价者评级","B6_full":"B6 全部特征","B7_interactions":"B7 B4＋预设交互（探索）","A1_without_competition":"完整模型去掉竞争组","A2_without_dynamics":"完整模型去掉时间组","A3_without_reputation":"完整模型去掉评级组"}
def md(df,cols=None):
    if cols:df=df[cols]
    def text(v):
        if pd.isna(v):return ""
        if isinstance(v,float):return f"{v:.3f}"
        return str(v).replace("|","/")
    return "| "+" | ".join(map(str,df.columns))+" |\n| "+" | ".join(["---"]*len(df.columns))+" |\n"+"\n".join("| "+" | ".join(text(v) for v in r)+" |" for r in df.itertuples(index=False,name=None))+"\n"
def main():
    m=pd.read_csv(OUT/"metrics_summary.csv")
    c=pd.read_csv(OUT/"paired_comparisons.csv")
    a=pd.read_csv(OUT/"single_feature_associations.csv")
    tail=pd.read_csv(OUT/"tail_sensitivity.csv")
    audit=json.loads((OUT/"audit.json").read_text(encoding="utf-8"))
    sections=["# Online Auctions 特征实验：B0—B6、消融与交互\n",
    "分析日期：2026-10-20。主预测时点为拍卖进行到50%；25%和75%为补充。所有结果保存于ftec5910_v2。\n",
    "## 1. 研究问题与实验边界\n",
    "研究问题：已知起拍条件与截点前价格信息后，竞争结构、时间变化及竞价者评级是否提供最终成交价的增量预测价值？分析的是拍卖进行中的条件价格预测，不是商品上市前估值或未来月份的市场价格预测，也不提供因果识别。\n",
    "原始Online Auctions auction.csv包含10,681次出价、628场拍卖。排除1场起拍价前后不一致的PDA拍卖3019271858，剩627场：Xbox149、PDA342、手表136；原文件未修改。PDA的2条记录出价超过最终价格的情况保留并披露。竞价者ID及评级缺失按规则处理，不将缺失ID视作同一个人。\n",
    "## 2. 实验设计\n",
    "- 每场拍卖在每个截点只有一行，输入只使用bidtime不晚于截点的记录。无出价的拍卖仍保留。\n- 每类商品分别建模，按拍卖时长分层做5折、重复3次。所有特征组和三个截点共用同一套拍卖折划分。\n- 填补缺失、标准化与模型拟合均在训练折完成；不使用拍卖ID、竞价者ID作为模型输入。\n- 固定Ridge alpha=10；浅层随机森林100棵树、最大深度5、最小叶样本5。未调参，因此本次不需要内部调参验证；如果以后调参，必须放在外层训练折内部。\n- 主要指标为MAE，另给RMSE、R²；不同商品不直接比较MAE的绝对大小。\n- 数据无真实日历开始时间，无法声称完成未来时间段验证。相同竞价者可能跨拍卖出现，未做竞价者隔离；已查看过数据，无未接触的最终测试。\n",
    "## 3. 特征定义与修正\n",
    "- B1：起拍价、拍卖时长。B2：再加入截点前最高记录出价、相对起拍价涨幅、是否已有出价。\n- B3：再加入出价次数、独立竞价人数、人均出价次数、竞价者出价占比HHI、缺失ID出价数。\n- B4：再加入首次出价时间、最近出价间隔、近期出价数/参与人数/真正新增人数/新增占比、近期出价占比、HHI变化和近期最高价变化。\n- B5：再加入已识别独立竞价者的反馈评级中位数、低评级占比及评级缺失占比。\n- B6：B2加上B3/B4/B5的附加组。消融实验分别去掉一个附加组。\n- B7为有限探索：B4加上真正新增人数×累计涨幅、HHI变化×累计涨幅；不是搜索任意特征组合。\n",
    "**重要修正：**旧脚本的recent_new_bidders实际只是近期活跃人数。本实验真正新增人数=近期出现、但窗口开始前从未出现的已知竞价者数；近期活跃人数另列。旧结果未覆盖，本报告使用新实验结果。近期窗口为总拍卖时长的10%，HHI变化为近期与前一相邻窗口之差。空窗口HHI定义为0，需结合是否有出价和人数解释。\n",
    "bid是记录出价，不保证等于当时公开显示的站立价格；评级可能是采集时快照，历史可得性仍未确认。低评级不能作为作弊标签。\n",
    "## 4. 单个特征：先理解，不按结果筛选\n",
    "Spearman相关为未控制其他因素的描述性关系，对每类商品和截点内的候选特征做BH校正。此步骤不用于挑选模型特征；强相关不等于对B2有增量。\n"]
    for key,name in NAMES.items():
        item={"xbox":"Xbox game console","pda":"Palm Pilot M515 PDA","watch":"Cartier wristwatch"}[key]
        t=a.loc[a.item.eq(item)&a.stage.eq(.5)&a.feature.isin(["current_max_bid","recent_new_bidders","hhi_change","bidder_hhi","rating_median"]),["feature","n","spearman","bh_q_value"]]
        sections += ["### "+name+"：50%截点\n",md(t)]
    sections+=["## 5. 按B0—B6顺序：50%主实验\n","下表为3次重复的整折外预测MAE平均值，不把3次重复当作3倍独立样本。B3—B6主要比较B2，同时检查是否优于B0基线。\n"]
    for key,name in NAMES.items():
        t=m.loc[m.item_key.eq(key)&m.stage.eq(.5)&m.experiment.str.startswith("B")].pivot(index="experiment",columns="model",values="mae")
        t=t.reindex(["B0_median","B0_current_bid","B1_opening","B2_price","B3_competition","B4_dynamics","B5_reputation","B6_full","B7_interactions"])
        t=t.rename(index=LABELS).reset_index().rename(columns={"experiment":"实验"})
        sections += ["### "+name+"\n",md(t)]
        for model in ["ridge","random_forest"]:
            b=c.loc[c.item_key.eq(key)&c.stage.eq(.5)&c.model.eq(model)&c.reference.eq("B2_price")&c.new.isin(["B3_competition","B4_dynamics","B5_reputation","B6_full","B7_interactions"])].copy()
            b["new"]=b["new"].map(LABELS)
            sections+=["相对B2的配对比较："+model+"（正数表示改善）\n",md(b[["new","mae_gain","gain_percent","ci_low","ci_high","positive_repeats"]])]
    sections+=["## 6. 去掉单组特征：消融检查\n","正改善表示完整B6优于去掉该组的模型；如果加组与去组结论不一致，说明可能存在冗余、交互或估计不稳定。\n"]
    t=c.loc[c.stage.eq(.5)&c.reference.str.startswith("A")].copy()
    t["item_key"]=t.item_key.map(NAMES)
    t["reference"]=t.reference.map(LABELS)
    sections+=[md(t[["item_key","model","reference","mae_gain","gain_percent","ci_low","ci_high"]])]
    sections+=["## 7. 不同时点及极高价敏感性\n"]
    t=m.loc[m.model.isin(["ridge","random_forest"])&m.experiment.isin(["B1_opening","B2_price","B4_dynamics","B6_full"])].copy()
    t["item_key"]=t.item_key.map(NAMES)
    sections+=[md(t[["item_key","stage","model","experiment","mae","rmse"]])]
    sections+=["尾部检查不重训、不删除原始数据：每个训练折计算成交价95%分位点，再把测试误差分成不超过该阈值与超过阈值两组。这样能看改善是否只由高价拍卖带动；不能据此事后选择删除异常点。表中记录数含3次重复，不是独立样本数。\n"]
    t=tail.loc[tail.stage.eq(.5)&tail.experiment.isin(["B0_median","B2_price","B4_dynamics","B6_full"])].copy()
    t["item_key"]=t.item_key.map(NAMES)
    sections+=[md(t[["item_key","model","experiment","stratum","n_repeat_records","mae"]])]
    sections+=["## 8. 不确定性与解释限制\n",
    "配对区间按同一拍卖在3次重复中的平均绝对误差差值，进行2,000次拍卖重采样；以已拟合交叉验证模型为条件，未覆盖全部重训练不确定性、训练折重叠、跨拍卖用户关系和模型/特征筛选，也未对全部模型比较作多重校正。不能将区间不跨0直接当作确认性显著发现。单特征BH校正与模型比较区间是两类不同分析。\n",
    "小样本下，不超过B2不代表这些信号在所有市场无效；可能是缺少库存、成色、商品配置、卖家和外部热度，或样本不足。手表价格26—5400，异质性较大；PDA价格相对集中。不要把收益最高的一个模型、商品或截点当作普遍结论。\n",
    "## 9. 文件与复现\n",
    "- 统一脚本：scripts/run_auction_feature_experiments.py。先prepare，再分别运行xbox/pda/watch，最后finish。\n- 汇总脚本：scripts/summarize_auction_feature_experiments.py。\n- Excel：outputs/auction_feature_experiments/auction_feature_experiments.xlsx。\n- prefix_features.csv与experiment_features.csv记录特征值及实验输入；fold_assignments.csv固定拍卖划分；每商品predictions.csv保留折外预测；audit.json记录数据SHA256、软件版本与参数。\n- *_single_features.png为主截点散点图；*_B_sequence.png为按实验顺序的模型比较图。\n"]
    # A compact decision table follows all results, with no significance overclaim.
    decisions=[]
    for key in NAMES:
        for family in ["ridge","random_forest"]:
            s=c.loc[c.item_key.eq(key)&c.stage.eq(.5)&c.model.eq(family)&c.reference.eq("B2_price")&c.new.isin(["B3_competition","B4_dynamics","B5_reputation","B6_full"])]
            for _,r in s.iterrows():
                decisions.append({"商品":NAMES[key],"模型":family,"实验":LABELS[r["new"]],"MAE改善百分比":r.gain_percent,"三个重复均改善":r.positive_repeats==3,"探索区间下界大于0":r.ci_low>0})
    decisions=pd.DataFrame(decisions)
    decisions.to_csv(OUT/"primary_decisions_zh.csv",index=False,encoding="utf-8-sig")
    sections.insert(1, '## 结果解读：按实验顺序\n\n**主要结论：已观察价格是可靠的基准信息，附加行为组尚未显示跨商品和模型稳定的增益。** B4时间变化组存在小幅线索，但当前证据不够确认；不能把“所有结果均为零”与“发现了稳定新信号”二者混为一谈。\n\n1. **B0/B1：建立基准。** Xbox与手表的仅起拍条件模型偏弱；PDA随机森林仅靠起拍条件即可优于全局中位数。B2以后始终既比较同模型B2，也检查B0，避免只超过弱基线。\n2. **B2：已观察价格带来主要收益。** 50%截点Ridge中，Xbox MAE由39.54降至31.19（21.12%）；手表由509.28降至212.75（58.23%）；PDA由17.39降至16.99（2.30%）。这是价格发现过程的预测价值，不是新行为因素的因果效应。PDA随机森林的B2为14.83，优于Ridge；模型优势随商品不同。\n3. **B3：累计竞争结构没有稳定额外收益。** PDA的Ridge仅改善0.67%，另外商品或模型不一致；不能把出价数或人数单独认定为高价值信号。\n4. **B4：最值得保留的探索线索是时间变化组。** Xbox随机森林从34.25降至33.35（2.64%）；手表Ridge从212.75降至208.48（2.01%）、随机森林从249.89降至245.32（1.83%），这些方向在3个重复中一致。但是对应配对区间仍跨0；Xbox的Ridge和PDA未改善，且手表Ridge的RMSE略变差。只能称候选，不能写成已经证明有效。\n5. **B5：反馈评级没有稳定收益。** 评级字段的历史可得性也不清楚，不支持将其作为主要发现。\n6. **B6：全部加入没有普遍胜出。** Xbox随机森林改善2.16%，但同商品Ridge变差；PDA、手表两模型均不优于各自B2。更多字段不等于更多可泛化信息。\n7. **消融：部分支持B4线索，但仍不确定。** 在Xbox与手表随机森林中，完整模型加入时间组有小幅改善，区间仍跨0；其他组与其他模型不一致。因此竞争和评级并非不可替代的信息来源。\n8. **B7：预设交互没有超过B4。** 在50%主截点的三个商品、两个模型中，加入这两个手工交互的MAE点估计均比B4更差；不继续事后穷举交互来挑最好结果。\n9. **敏感性：B4的小幅改善并非单纯由最高价样本推动。** Xbox随机森林普通价格组MAE从29.69降至28.69，但高价组从122.41升至123.34；手表随机森林普通组204.80降至198.73，高价组940.68升至959.11。收益在不同价格区域不一致，进一步限制普遍结论。\n\n以上结论是组级别的，不能证明“真正新增人数”或“HHI变化”中哪一列单独贡献了收益。附加特征相对B2的所有探索性配对区间，在25%、50%、75%三个时点中均未出现下界大于0的改善；这不证明真实效果必为零。\n\n### 可用于报告的段落\n\n> 本研究在Online Auctions数据上，以拍卖进度50%为主要预测截点，对三类商品分别采用固定参数Ridge及浅层随机森林，执行三次重复的五折拍卖层面交叉验证。通过B0—B6递增实验、单组消融及少量预设交互，发现截点前已记录价格是预测最终成交价的主要信息来源。竞争人数与集中度、竞价者评级尚未提供稳定增量；时间变化组在Xbox随机森林及手表模型中呈现约2%的MAE改善，但配对区间包含零，且跨模型和商品不一致。因此，时间变化可作为后续研究候选，当前不足以宣称发现了稳定的新价格驱动因素。\n\n')
    (ROOT/"docs/auction_feature_experiments_zh.md").write_text("\n".join(sections),encoding="utf-8")
    quick=["# 拍卖特征实验：结果速读\n","主实验：50%时点，3次重复5折，商品分别建模。结果是探索性，不能作因果解释。\n"]
    for key,name in NAMES.items():
        quick+=["## "+name+"\n"]
        for family in ["ridge","random_forest"]:
            g=m.loc[m.item_key.eq(key)&m.stage.eq(.5)&m.model.eq(family)].set_index("experiment")
            b2=g.loc["B2_price","mae"];b6=g.loc["B6_full","mae"];b4=g.loc["B4_dynamics","mae"]
            baseline=m.loc[m.item_key.eq(key)&m.stage.eq(.5)&m.model.eq("baseline"),"mae"].min()
            quick+=[f"- {family}：B2 MAE={b2:.2f}；B4={b4:.2f}；B6={b6:.2f}。B6相对B2改善={(b2-b6)/b2*100:.2f}%；两个B0基线中较低MAE={baseline:.2f}。\n"]
    quick+=["## 如何使用\n","正改善仅表示本次误差更小，需结合重复方向、区间、消融和补充截点；不要以局部最优值宣称普遍有效。全部实验与数字见 [完整报告](auction_feature_experiments_zh.md)。\n"]
    quick.insert(1, '## 结果解读：按实验顺序\n\n**主要结论：已观察价格是可靠的基准信息，附加行为组尚未显示跨商品和模型稳定的增益。** B4时间变化组存在小幅线索，但当前证据不够确认；不能把“所有结果均为零”与“发现了稳定新信号”二者混为一谈。\n\n1. **B0/B1：建立基准。** Xbox与手表的仅起拍条件模型偏弱；PDA随机森林仅靠起拍条件即可优于全局中位数。B2以后始终既比较同模型B2，也检查B0，避免只超过弱基线。\n2. **B2：已观察价格带来主要收益。** 50%截点Ridge中，Xbox MAE由39.54降至31.19（21.12%）；手表由509.28降至212.75（58.23%）；PDA由17.39降至16.99（2.30%）。这是价格发现过程的预测价值，不是新行为因素的因果效应。PDA随机森林的B2为14.83，优于Ridge；模型优势随商品不同。\n3. **B3：累计竞争结构没有稳定额外收益。** PDA的Ridge仅改善0.67%，另外商品或模型不一致；不能把出价数或人数单独认定为高价值信号。\n4. **B4：最值得保留的探索线索是时间变化组。** Xbox随机森林从34.25降至33.35（2.64%）；手表Ridge从212.75降至208.48（2.01%）、随机森林从249.89降至245.32（1.83%），这些方向在3个重复中一致。但是对应配对区间仍跨0；Xbox的Ridge和PDA未改善，且手表Ridge的RMSE略变差。只能称候选，不能写成已经证明有效。\n5. **B5：反馈评级没有稳定收益。** 评级字段的历史可得性也不清楚，不支持将其作为主要发现。\n6. **B6：全部加入没有普遍胜出。** Xbox随机森林改善2.16%，但同商品Ridge变差；PDA、手表两模型均不优于各自B2。更多字段不等于更多可泛化信息。\n7. **消融：部分支持B4线索，但仍不确定。** 在Xbox与手表随机森林中，完整模型加入时间组有小幅改善，区间仍跨0；其他组与其他模型不一致。因此竞争和评级并非不可替代的信息来源。\n8. **B7：预设交互没有超过B4。** 在50%主截点的三个商品、两个模型中，加入这两个手工交互的MAE点估计均比B4更差；不继续事后穷举交互来挑最好结果。\n9. **敏感性：B4的小幅改善并非单纯由最高价样本推动。** Xbox随机森林普通价格组MAE从29.69降至28.69，但高价组从122.41升至123.34；手表随机森林普通组204.80降至198.73，高价组940.68升至959.11。收益在不同价格区域不一致，进一步限制普遍结论。\n\n以上结论是组级别的，不能证明“真正新增人数”或“HHI变化”中哪一列单独贡献了收益。附加特征相对B2的所有探索性配对区间，在25%、50%、75%三个时点中均未出现下界大于0的改善；这不证明真实效果必为零。\n\n### 可用于报告的段落\n\n> 本研究在Online Auctions数据上，以拍卖进度50%为主要预测截点，对三类商品分别采用固定参数Ridge及浅层随机森林，执行三次重复的五折拍卖层面交叉验证。通过B0—B6递增实验、单组消融及少量预设交互，发现截点前已记录价格是预测最终成交价的主要信息来源。竞争人数与集中度、竞价者评级尚未提供稳定增量；时间变化组在Xbox随机森林及手表模型中呈现约2%的MAE改善，但配对区间包含零，且跨模型和商品不一致。因此，时间变化可作为后续研究候选，当前不足以宣称发现了稳定的新价格驱动因素。\n\n')
    (ROOT/"docs/auction_feature_experiments_quick_zh.md").write_text("\n".join(quick),encoding="utf-8")
    register=ROOT/"docs/data_candidates.csv"
    if register.exists():
        d=pd.read_csv(register,encoding="utf-8-sig",keep_default_na=False)
        mask=d.iloc[:,0].eq("D005")
        d.loc[mask,d.columns[4]]="2026-10-20: full B0-B6 / ablation / limited-interaction experiments completed; true new-bidder definition corrected"
        d.loc[mask,d.columns[12]]="See docs/auction_feature_experiments_zh.md; primary cutoff 50%; matched 3x5 auction CV; fixed Ridge and shallow forest"
        d.to_csv(register,index=False,encoding="utf-8-sig")
    book=load_workbook(OUT/"auction_feature_experiments.xlsx",read_only=True)
    assert len(book.sheetnames)==8
    assert len(m)==3*3*(2+2*10)
    print(decisions.to_string(index=False))
    print("Chinese full/quick reports, decision CSV and source register saved.")
if __name__=="__main__":main()
